"""
Generate a full Excel report: PO balances, product-level breakdown, and invoice history.
Run: python3 generate_excel_report.py
Output: data/reports/PO_Balance_Report.xlsx
"""
import sqlite3
import os
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DB_PATH = os.path.join(os.path.dirname(__file__), "data", "po_tracker.db")
OUT_DIR = os.path.join(os.path.dirname(__file__), "data", "reports")
OUT_PATH = os.path.join(OUT_DIR, "PO_Balance_Report.xlsx")

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14)
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STATUS_FILL = {
    "COMPLETED": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
    "IN PROGRESS": PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid"),
    "NOT STARTED": PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid"),
    "OVERDUE": PatternFill(start_color="FECACA", end_color="FECACA", fill_type="solid"),
}


def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER


def _autosize(ws):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(max(length + 3, 12), 45)


def _determine_status(ordered, dispatched, due_date_str):
    if ordered <= 0:
        return "NOT STARTED"
    pct = dispatched / ordered * 100
    if pct >= 100:
        return "COMPLETED"
    if due_date_str:
        try:
            due = date.fromisoformat(due_date_str)
            if due < date.today() and dispatched < ordered:
                return "OVERDUE"
        except ValueError:
            pass
    if pct == 0:
        return "NOT STARTED"
    return "IN PROGRESS"


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    cur = conn.cursor()

    pos = cur.execute("SELECT * FROM PURCHASE_ORDERS ORDER BY PODate").fetchall()

    wb = Workbook()

    # ---------- Sheet 1: PO Summary ----------
    ws1 = wb.active
    ws1.title = "PO Summary"
    ws1["A1"] = "Purchase Order — Balance Summary"
    ws1["A1"].font = TITLE_FONT
    ws1["A2"] = f"Generated: {date.today().isoformat()}"

    headers = ["PO Number", "PO Date", "Due Date", "Status", "Total Ordered", "Total Dispatched", "Balance Pending", "% Complete"]
    ws1.append([])
    ws1.append(headers)
    _style_header(ws1, 4, len(headers))

    po_summaries = []
    for po in pos:
        items = cur.execute("SELECT * FROM PURCHASE_ORDER_ITEMS WHERE POID=?", (po["POID"],)).fetchall()
        total_ordered = 0
        total_dispatched = 0
        for item in items:
            dispatched = cur.execute(
                """SELECT COALESCE(SUM(ii.DispatchedQty),0) AS total
                   FROM INVOICE_ITEMS ii JOIN INVOICES inv ON ii.InvoiceID = inv.InvoiceID
                   WHERE ii.ProductCode=? AND inv.PONumber=?""",
                (item["ProductCode"], po["PONumber"]),
            ).fetchone()["total"]
            total_ordered += item["OrderedQty"]
            total_dispatched += dispatched
        pending = max(0, total_ordered - total_dispatched)
        status = _determine_status(total_ordered, total_dispatched, po["DueDate"])
        pct = (total_dispatched / total_ordered * 100) if total_ordered > 0 else 0
        po_summaries.append({
            "po": po, "total_ordered": total_ordered, "total_dispatched": total_dispatched,
            "pending": pending, "status": status, "pct": pct,
        })

        row = [po["PONumber"], po["PODate"] or "—", po["DueDate"] or "—", status,
               total_ordered, total_dispatched, pending, round(pct, 1)]
        ws1.append(row)
        r = ws1.max_row
        for c in range(1, len(headers) + 1):
            ws1.cell(row=r, column=c).border = BORDER
        fill = STATUS_FILL.get(status)
        if fill:
            for c in range(1, len(headers) + 1):
                ws1.cell(row=r, column=c).fill = fill

    _autosize(ws1)
    ws1.freeze_panes = "A5"

    # ---------- Sheet 2: Product-level Breakdown ----------
    ws2 = wb.create_sheet("Product Breakdown")
    headers2 = ["PO Number", "PO Date", "Product Code", "Product Name", "Ordered Qty", "Dispatched Qty", "Pending Qty", "Status"]
    ws2.append(headers2)
    _style_header(ws2, 1, len(headers2))

    for s in po_summaries:
        po = s["po"]
        items = cur.execute("SELECT * FROM PURCHASE_ORDER_ITEMS WHERE POID=?", (po["POID"],)).fetchall()
        for item in items:
            dispatched = cur.execute(
                """SELECT COALESCE(SUM(ii.DispatchedQty),0) AS total
                   FROM INVOICE_ITEMS ii JOIN INVOICES inv ON ii.InvoiceID = inv.InvoiceID
                   WHERE ii.ProductCode=? AND inv.PONumber=?""",
                (item["ProductCode"], po["PONumber"]),
            ).fetchone()["total"]
            pending = max(0, item["OrderedQty"] - dispatched)
            status = _determine_status(item["OrderedQty"], dispatched, po["DueDate"])
            ws2.append([po["PONumber"], po["PODate"] or "—", item["ProductCode"], item["ProductName"],
                        item["OrderedQty"], dispatched, pending, status])
            r = ws2.max_row
            for c in range(1, len(headers2) + 1):
                ws2.cell(row=r, column=c).border = BORDER
            fill = STATUS_FILL.get(status)
            if fill:
                for c in range(1, len(headers2) + 1):
                    ws2.cell(row=r, column=c).fill = fill

    _autosize(ws2)
    ws2.freeze_panes = "A2"

    # ---------- Sheet 3: Invoices Given ----------
    ws3 = wb.create_sheet("Invoices Given")
    headers3 = ["Invoice Number", "Invoice Date", "PO Number", "Product Code", "Product Name", "Dispatched Qty"]
    ws3.append(headers3)
    _style_header(ws3, 1, len(headers3))

    invoices = cur.execute("SELECT * FROM INVOICES ORDER BY InvoiceDate").fetchall()
    for inv in invoices:
        items = cur.execute("SELECT * FROM INVOICE_ITEMS WHERE InvoiceID=?", (inv["InvoiceID"],)).fetchall()
        for item in items:
            ws3.append([inv["InvoiceNumber"], inv["InvoiceDate"] or "—", inv["PONumber"] or "—",
                        item["ProductCode"], item["ProductName"], item["DispatchedQty"]])
            r = ws3.max_row
            for c in range(1, len(headers3) + 1):
                ws3.cell(row=r, column=c).border = BORDER

    _autosize(ws3)
    ws3.freeze_panes = "A2"

    # ---------- Sheet 4: Pending Action Items (what to dispatch next) ----------
    ws4 = wb.create_sheet("Pending To Give")
    headers4 = ["PO Number", "Due Date", "Product Code", "Product Name", "Pending Qty", "Status"]
    ws4.append(headers4)
    _style_header(ws4, 1, len(headers4))

    for s in po_summaries:
        po = s["po"]
        if s["status"] == "COMPLETED":
            continue
        items = cur.execute("SELECT * FROM PURCHASE_ORDER_ITEMS WHERE POID=?", (po["POID"],)).fetchall()
        for item in items:
            dispatched = cur.execute(
                """SELECT COALESCE(SUM(ii.DispatchedQty),0) AS total
                   FROM INVOICE_ITEMS ii JOIN INVOICES inv ON ii.InvoiceID = inv.InvoiceID
                   WHERE ii.ProductCode=? AND inv.PONumber=?""",
                (item["ProductCode"], po["PONumber"]),
            ).fetchone()["total"]
            pending = max(0, item["OrderedQty"] - dispatched)
            if pending > 0:
                status = _determine_status(item["OrderedQty"], dispatched, po["DueDate"])
                ws4.append([po["PONumber"], po["DueDate"] or "—", item["ProductCode"], item["ProductName"], pending, status])
                r = ws4.max_row
                for c in range(1, len(headers4) + 1):
                    ws4.cell(row=r, column=c).border = BORDER
                fill = STATUS_FILL.get(status)
                if fill:
                    for c in range(1, len(headers4) + 1):
                        ws4.cell(row=r, column=c).fill = fill

    _autosize(ws4)
    ws4.freeze_panes = "A2"

    # ---------- Sheet 5: Balance To Distribute (clean, qty only) ----------
    ws5 = wb.create_sheet("Balance To Distribute")
    headers5 = ["PO Number", "Product Name", "Balance Qty"]
    ws5.append(headers5)
    _style_header(ws5, 1, len(headers5))

    for s in po_summaries:
        po = s["po"]
        if s["status"] == "COMPLETED":
            continue
        items = cur.execute("SELECT * FROM PURCHASE_ORDER_ITEMS WHERE POID=?", (po["POID"],)).fetchall()
        for item in items:
            dispatched = cur.execute(
                """SELECT COALESCE(SUM(ii.DispatchedQty),0) AS total
                   FROM INVOICE_ITEMS ii JOIN INVOICES inv ON ii.InvoiceID = inv.InvoiceID
                   WHERE ii.ProductCode=? AND inv.PONumber=?""",
                (item["ProductCode"], po["PONumber"]),
            ).fetchone()["total"]
            pending = max(0, item["OrderedQty"] - dispatched)
            if pending > 0:
                ws5.append([po["PONumber"], item["ProductName"], pending])
                r = ws5.max_row
                for c in range(1, len(headers5) + 1):
                    ws5.cell(row=r, column=c).border = BORDER

    _autosize(ws5)
    ws5.freeze_panes = "A2"

    wb.save(OUT_PATH)
    conn.close()
    print(f"Report saved to: {OUT_PATH}")


if __name__ == "__main__":
    main()
